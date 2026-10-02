"""Derive the reviewed deployed library LIB-FAITHFUL-P-v2 (8 rules).

Supports: the library used by CP-GTR and CP-GTR-Detect in the compliance
evaluation, and the rule-fidelity review. It re-applies the membership rule of
LIB-FAITHFUL-P (supported_by_span in {yes, partial} and guard_correct != no) to
the reviewer's fresh verdicts on the 13 rules whose templates were rewritten.
Where the reviewer marked verdict_still_holds="changed", the fresh
classification is the one stated in his notes (FRESH_CLASSIFICATION below,
keyed by content hash); otherwise the original classification applies.
guard_correct is "n-a" for all 13 rules, so membership is decided by
supported_by_span alone.

A new entry "LIB-FAITHFUL-P-v2" is written into analysis/phase2_library.json.
The original 13-rule "LIB-FAITHFUL-P" entry is not modified, so both coexist
and downstream code must name the one it wants.

Run:
    python analysis/phase2_build_faithful_p_v2.py
Order: after phase2_build_fidelity_recheck_xlsx.py and the reviewer's filled
workbook; before phase2_generation_run.py, which asserts the resulting hash.
Same family as phase2_build_faithful_v2.py (strict variant) and
phase2_build_lib_full_v2.py.

Inputs : docs/fidelity_recheck_filled__Prof_Ali.xlsx and
         docs/fidelity_audit_Ali_filled2.csv (both human annotation data, not
         part of the public release), analysis/phase2_library.json.
Outputs: analysis/phase2_library.json (adds the LIB-FAITHFUL-P-v2 entry).
"""
import json
import sys
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT))

from analysis.phase2_harness import _rows_sha256, PHASE2_LIBRARY_PATH

FILLED_XLSX = REPO_ROOT / "docs" / "fidelity_recheck_filled__Prof_Ali.xlsx"

wb = openpyxl.load_workbook(FILLED_XLSX)
ws = wb["Fidelity Recheck"]

verdicts = {}   # content_hash -> (verdict_still_holds, notes)
for r in range(2, ws.max_row + 1):
    rh = ws[f"A{r}"].value
    if not rh:
        continue
    verdicts[rh] = (ws[f"G{r}"].value, ws[f"H{r}"].value or "")

assert len(verdicts) == 13, len(verdicts)

# Fresh classification for every "changed" row, transcribed from the
# reviewer's notes. The workbook dropdown only offers yes/no/changed, and
# "changed" means a fresh judgement was made, not that it is itself a
# classification.
FRESH_CLASSIFICATION = {
    # content_hash: "yes" | "partial" | "no", per the reviewer's notes.
    # OperatorCollectsPersistentIdentifierInternalUse_NoNotice
    "3e62a311cc627fa1a3ed097e14fffdef25bd758e6d2b70645b524f4949122af6": "no",
    # Provision3 (uk_aadc2)
    "e477c6fddc4b07408eb10410afc84324d517c0e099c6a34910a53edf644d116d": "no",
    # Provision4 (coppa)
    "1a2490f719af18c502c5c8628292897d082710a143cd6192a1311adacf349ceb": "no",
    # StatuteOfLimitations_AdminAction
    "9db8f19ab63934788a6bace3665dae54ed661e28a6cc9d6a5639ae271776985e": "no",
    # Provision2 (coppa)
    "ddd030d02f418324db87eab951301a08b72d601557d88423bfec5c8c03a655ec": "no",
    # Provision1 (uk_aadc2, default sharing)
    "6b907952ebe4b4f225932dee0334b9618da9ac394c6e9899855ea9d971c32abb": "partial",
}

lib_data = json.loads(PHASE2_LIBRARY_PATH.read_text(encoding="utf-8"))
old_entry = lib_data["LIB-FAITHFUL-P"]
old_rules = old_entry["rules"]
assert len(old_rules) == 13

kept, dropped = [], []
disposition = {}
for row in old_rules:
    ch = row["content_hash"]
    verdict, notes = verdicts[ch]
    if verdict == "changed":
        fresh = FRESH_CLASSIFICATION[ch]
    else:
        # verdict_still_holds == "yes" means the original supported_by_span
        # classification (yes or partial) still applies; it does not mean the
        # value is literally the string "yes".
        with open(REPO_ROOT / "docs" / "fidelity_audit_Ali_filled2.csv", encoding="utf-8-sig") as f:
            import csv
            orig = {r["rule_id"]: r["supported_by_span"] for r in csv.DictReader(f)}
        fresh = orig[ch]
    qualifies = fresh in ("yes", "partial")
    disposition[row["name"]] = {"content_hash": ch, "fresh_classification": fresh,
                                 "verdict_still_holds": verdict, "qualifies": qualifies,
                                 "notes": notes}
    (kept if qualifies else dropped).append(row)

print(f"Kept: {len(kept)}  Dropped: {len(dropped)}")
for name, d in disposition.items():
    print(f"  {'KEEP' if d['qualifies'] else 'DROP':5s} {name:55s} fresh={d['fresh_classification']}")

new_sha256 = _rows_sha256(kept)

new_entry = {
    "library_id": "LIB-FAITHFUL-P-v2",
    "sha256": new_sha256,
    "n_rules": len(kept),
    "membership_source": "EXTERNAL EXPERT AUDIT (Prof. Ali, independent judicial expert), "
                          "RE-VERIFIED AGAINST TASK I'S REWRITTEN TEMPLATES -- "
                          "docs/fidelity_recheck_filled__Prof_Ali.xlsx (analysis/fidelity_recheck.csv, "
                          "seed=7). Supersedes LIB-FAITHFUL-P (13 rules) for any use where the "
                          "REWRITTEN (notice-sentence) templates are what gets deployed -- the "
                          "original LIB-FAITHFUL-P audit verified the OLD drafter-instruction "
                          "template strings, which are no longer what CP-GTR emits.",
    "definition": "Same rule as LIB-FAITHFUL-P: supported_by_span in (yes, partial) AND "
                  "guard_correct != no -- applied to Prof. Ali's FRESH verdicts (6 of 13 rules "
                  "re-judged; 5 dropped to 'no' support for the new template's claim, 1 to "
                  "'partial'). guard_correct is 'n-a' for all 13 rules, so it never gates here.",
    "superseded": "LIB-FAITHFUL-P",
    "dropped_from_LIB-FAITHFUL-P": [
        {"name": name, "content_hash": d["content_hash"], "fresh_classification": d["fresh_classification"],
         "notes": d["notes"]}
        for name, d in disposition.items() if not d["qualifies"]
    ],
    "rules": kept,
}

lib_data["LIB-FAITHFUL-P-v2"] = new_entry
PHASE2_LIBRARY_PATH.write_text(json.dumps(lib_data, indent=2, ensure_ascii=False), encoding="utf-8")
print(f"\nWritten: LIB-FAITHFUL-P-v2 (n_rules={len(kept)}, sha256={new_sha256})")
print(f"Original LIB-FAITHFUL-P entry (13 rules, sha256={old_entry['sha256'][:16]}...) left UNTOUCHED, both coexist.")
