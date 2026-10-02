"""Build the reviewer workbook for the fidelity re-check.

Supports: the rule-fidelity review (fidelity table and the derivation of
LIB-FAITHFUL-P-v2). It turns analysis/fidelity_recheck.csv into the protected
spreadsheet a reviewer fills in, using the same conventions as the codebook
workbook: Arial font, navy headers for reference columns and orange headers
for reviewer columns, wrapped long columns, a frozen header row, locked
reference columns and unlocked reviewer columns, sheet protection without a
password, and an "About" sheet with provenance. Row order is read verbatim
from the CSV (already shuffled by phase2_taskN_build_recheck.py) and is not
re-shuffled or sorted here.

Run:
    python analysis/phase2_build_fidelity_recheck_xlsx.py
Order: after phase2_taskN_build_recheck.py; the filled workbook is the input of
phase2_build_faithful_p_v2.py and phase2_build_faithful_v2.py.

Inputs : analysis/fidelity_recheck.csv,
         analysis/realization_templates_corrections.json (hash check only).
Outputs: schema/fidelity_recheck.xlsx.
"""
import csv
import hashlib
import math
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.protection import SheetProtection
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parent.parent
CSV_PATH = REPO_ROOT / "analysis" / "fidelity_recheck.csv"
CORRECTIONS_PATH = REPO_ROOT / "analysis" / "realization_templates_corrections.json"
OUT_PATH = REPO_ROOT / "schema" / "fidelity_recheck.xlsx"

NAVY = "1F4E78"
ORANGE = "E36C09"
WHITE = "FFFFFF"

with open(CSV_PATH, encoding="utf-8") as f:
    rows = list(csv.DictReader(f))
assert len(rows) == 13, len(rows)

csv_sha256 = hashlib.sha256(CSV_PATH.read_bytes()).hexdigest()
corrections_sha256 = hashlib.sha256(CORRECTIONS_PATH.read_bytes()).hexdigest()
assert corrections_sha256 == "e4121e77ac0b1fee01f652a5981e96b688dc744b8116eb5d144fcf8c3808cc9d"

COLUMNS = [
    # (header, csv_key, width, wrap, locked, section)
    ("rule_id", "rule_id", 68, False, True, "ref"),
    ("source_file", "source_file", 20, False, True, "ref"),
    ("source_provision_text", "source_provision_text", 95, True, True, "ref"),
    ("old_template", "old_template", 45, True, True, "ref"),
    ("new_template", "new_template", 45, True, True, "ref"),
    ("supported_by_span", "supported_by_span", 16, False, True, "ref"),
    ("verdict_still_holds", None, 18, False, False, "reviewer"),
    ("notes", None, 40, True, False, "reviewer"),
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Fidelity Recheck"

header_font = Font(name="Arial", size=11, bold=True, color=WHITE)
header_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
data_font = Font(name="Arial", size=10, bold=False)
data_align_wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)
data_align_nowrap = Alignment(horizontal="left", vertical="top", wrap_text=False)
thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

ref_fill = PatternFill("solid", fgColor=NAVY)
reviewer_fill = PatternFill("solid", fgColor=ORANGE)

# ---- header row ----
for i, (header, _key, width, _wrap, _locked, section) in enumerate(COLUMNS, start=1):
    col = get_column_letter(i)
    cell = ws[f"{col}1"]
    cell.value = header
    cell.font = header_font
    cell.alignment = header_align
    cell.fill = ref_fill if section == "ref" else reviewer_fill
    cell.protection = openpyxl.styles.Protection(locked=True)  # header always locked
    ws.column_dimensions[col].width = width
ws.row_dimensions[1].height = 22

# ---- data rows (verbatim CSV order, no re-shuffle or sort) ----
for r_idx, row in enumerate(rows, start=2):
    max_lines = 1
    for c_idx, (header, key, width, wrap, locked, section) in enumerate(COLUMNS, start=1):
        col = get_column_letter(c_idx)
        value = row[key] if key is not None else None
        cell = ws[f"{col}{r_idx}"]
        cell.value = value
        cell.font = data_font
        cell.alignment = data_align_wrap if wrap else data_align_nowrap
        cell.protection = openpyxl.styles.Protection(locked=locked)
        cell.border = border
        if wrap and value:
            # Excel column width units approximate characters at this font
            # size; estimate the wrapped line count so no cell is truncated.
            chars_per_line = max(10, width - 2)
            lines = math.ceil(len(str(value)) / chars_per_line)
            max_lines = max(max_lines, lines)
    ws.row_dimensions[r_idx].height = max(40, max_lines * 14 + 10)

ws.freeze_panes = "A2"

# ---- dropdown validation on verdict_still_holds (column G) ----
dv = DataValidation(type="list", formula1='"yes,no,changed"', allow_blank=True,
                     showDropDown=False)  # False shows the arrow (openpyxl quirk)
dv.error = "Choose one of: yes, no, changed"
dv.errorTitle = "Invalid entry"
dv.prompt = "yes / no / changed"
dv.promptTitle = "Verdict still holds?"
ws.add_data_validation(dv)
last_row = len(rows) + 1
dv.add(f"G2:G{last_row}")

# ---- sheet protection on, without a password (as in the codebook workbook) ----
ws.protection = SheetProtection(
    sheet=True, password=None,
    selectLockedCells=False, selectUnlockedCells=False,
    insertRows=True, insertHyperlinks=True, autoFilter=True,
    deleteColumns=True, insertColumns=True, pivotTables=True, deleteRows=True,
    formatCells=False, formatColumns=False, formatRows=False, sort=False,
)

# =================================================================== About
about = wb.create_sheet("About")
about.column_dimensions["A"].width = 22
about.column_dimensions["B"].width = 100

label_font = Font(name="Arial", size=10, bold=True)
value_font = Font(name="Arial", size=10, bold=False)
value_align = Alignment(horizontal="left", vertical="top", wrap_text=True)
label_align = Alignment(horizontal="left", vertical="top")

about_rows = [
    ("fidelity_recheck.csv SHA-256", csv_sha256),
    ("Shuffle seed", "7 (random.Random(7).shuffle -- distinct from the original "
                      "fidelity_audit.csv's seed=1; analysis/phase2_taskN_build_recheck.py)"),
    ("realization_templates_corrections.json SHA-256", corrections_sha256),
    ("Code tag", "phase2-templates-corrected"),
    ("What changed", "The new_template column states only what each rule's R "
                      "(the graph-transformation right-hand side) adds or removes "
                      "relative to L -- it carries no age or jurisdiction condition "
                      "the rule's own Phi guard does not enforce. See "
                      "analysis/realization_templates.md for the full old/new record "
                      "and per-rule notes on every rewrite."),
]
for i, (label, value) in enumerate(about_rows, start=1):
    a, b = about[f"A{i}"], about[f"B{i}"]
    a.value, b.value = label, value
    a.font, b.font = label_font, value_font
    a.alignment, b.alignment = label_align, value_align
    lines = math.ceil(len(str(value)) / 98) if value else 1
    about.row_dimensions[i].height = max(18, lines * 14 + 8)

wb.save(OUT_PATH)
print(f"Written: {OUT_PATH}")
print(f"Sheet 'Fidelity Recheck': {ws.max_row - 1} data rows, {ws.max_column} columns")
print(f"CSV SHA-256: {csv_sha256}")
