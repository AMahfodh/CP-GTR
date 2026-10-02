"""Build the annotation workbooks and the overlap sample.

Supports: the annotation protocol behind the compliance results: 353 unique
annotation units, one primary annotator who judges every unit, and two further
annotators who independently judge an 82-unit overlap sample (23%), used for
the inter-annotator agreement (kappa) results and for majority-vote
adjudication.

Reads the blind annotation sheet (353 units, shuffled with seed 2026) and
schema/codebook_v2.json, and writes two protected workbooks with identical
conventions: Arial font, navy headers for reference columns and orange headers
for judgement columns, wrapped long columns, frozen panes, locked reference
and not-applicable cells with unlocked judgement cells, and sheet protection
without a password. Each workbook has Practice, Annotation, Codes and About
sheets.

Blindness: the sheet reaching the workbooks carries only {row_id, text, A, J}.
Arm, library, prompt id, flag status, round-trip result, rules fired,
provenance and parse diagnostics are dropped upstream by
analysis/phase2_harness.py (build_scored_documents() and
build_blind_annotation_sheet()) and never written into a workbook. The source
arm of each unit is recomputed here from analysis/phase2_generation_pairs.json
(never from the CSV) only to stratify the overlap sample, and is not written to
either workbook.

Overlap sample: a stratified draw of 71 units (20%), balanced over
(source arm, age, jurisdiction) by largest-remainder allocation with seed 2026,
unioned with every CP-GTR-originated unit (12) not already drawn. The union
gives full coverage of the arm under test and proportional stratification
elsewhere; it is deterministic, so it consumes no further random draw. Both the
original draw and the final membership are recorded in overlap_sample.json.

Run:
    python analysis/phase2_build_annotation_workbooks.py
Order: after phase2_taskZ_annotation_sheet.py; before the annotators fill the
workbooks and before phase2_final_results.py.

Inputs : analysis/phase2_blind_annotation_sheet.csv (353 units; LLM-generated
         text), analysis/phase2_generation_pairs.json,
         analysis/phase2_generation_practice.json, schema/codebook_v2.json.
Outputs: annotation/annotator_primary.xlsx, annotation/annotator_overlap.xlsx
         (blank workbooks), annotation/overlap_sample.json.
"""
import csv
import hashlib
import json
import math
import random
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.protection import SheetProtection

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT))

from analysis.phase2_harness import build_scored_documents, text_hash, unique_units

CSV_PATH = REPO_ROOT / "analysis" / "phase2_blind_annotation_sheet.csv"
CODEBOOK_PATH = REPO_ROOT / "schema" / "codebook_v2.json"
PAIRS_PATH = REPO_ROOT / "analysis" / "phase2_generation_pairs.json"
PRACTICE_PATH = REPO_ROOT / "analysis" / "phase2_generation_practice.json"
OUT_DIR = REPO_ROOT / "annotation"
OUT_DIR.mkdir(exist_ok=True)
PRIMARY_PATH = OUT_DIR / "annotator_primary.xlsx"
OVERLAP_PATH = OUT_DIR / "annotator_overlap.xlsx"
OVERLAP_SAMPLE_JSON = OUT_DIR / "overlap_sample.json"

SEED = 2026
NAVY = "1F4E78"
ORANGE = "E36C09"
WHITE = "FFFFFF"
STRIPE = "F2F2F2"
PRACTICE_BANNER_FILL = "C00000"

# --------------------------------------------------------------------------- #
# Load source data
# --------------------------------------------------------------------------- #
with open(CSV_PATH, encoding="utf-8") as f:
    csv_rows = list(csv.DictReader(f))
assert len(csv_rows) == 353, len(csv_rows)

csv_sha256 = hashlib.sha256(CSV_PATH.read_bytes()).hexdigest()
codebook_bytes = CODEBOOK_PATH.read_bytes()
codebook_sha256_wholefile = hashlib.sha256(codebook_bytes).hexdigest()
codebook = json.loads(codebook_bytes)
# Hash recorded in the codebook itself; its basis is {codes, over_requirement}
# only, so it differs from the whole-file hash by design.
codebook_embedded_sha256 = codebook["sha256"]

codes = codebook["codes"] + [codebook["over_requirement"]]
code_ids = [c["code"] for c in codes]
assert len(code_ids) == 11, code_ids
assert code_ids[-1] == "OVER_REQUIREMENT"


def code_applies(code_entry, A, J):
    """applies_when gating: whether `code_entry` can apply to a document read by
    age A in jurisdiction J. OVER_REQUIREMENT is never gated; it applies to
    every document."""
    if code_entry["code"] == "OVER_REQUIREMENT":
        return True
    aw = code_entry["applies_when"]
    j_cond = aw["jurisdiction"]
    if j_cond != "any" and j_cond != J:
        return False
    age_cond = aw["age_condition"]
    if age_cond.startswith("any"):
        return True
    m = re.match(r"A\s*<\s*(\d+)", age_cond)
    assert m, f"unparseable age_condition: {age_cond!r}"
    return A < int(m.group(1))


# --------------------------------------------------------------------------- #
# Internal only: source arm per unit, used to stratify the overlap sample.
# Recomputed from phase2_generation_pairs.json (the six-arm source), never from
# the CSV, and written to neither workbook; it appears only in the balance
# report printed to the console and in overlap_sample.json's balance counts.
# --------------------------------------------------------------------------- #
pairs = json.loads(PAIRS_PATH.read_text(encoding="utf-8"))
scored_docs = build_scored_documents(pairs)
assert len(unique_units(scored_docs)) == 353

source_arm_by_unit = {}
for d in scored_docs:
    # The first arm to produce a unit is its source, as in inheritance_report().
    source_arm_by_unit.setdefault(d["unit"], d["arm"])


def row_unit(row):
    return (text_hash(row["text"]), int(row["A"]), row["J"])


row_source_arm = {}
for row in csv_rows:
    u = row_unit(row)
    assert u in source_arm_by_unit, f"CSV row_id={row['row_id']} has no matching scored-document unit"
    row_source_arm[row["row_id"]] = source_arm_by_unit[u]

# --------------------------------------------------------------------------- #
# Stratified 20% overlap sample (71 of 353), balanced across (source arm, A, J),
# largest-remainder allocation, seed 2026.
# --------------------------------------------------------------------------- #
TARGET_N = round(len(csv_rows) * 0.20)
assert TARGET_N == 71, TARGET_N

strata = defaultdict(list)
for row in csv_rows:
    key = (row_source_arm[row["row_id"]], int(row["A"]), row["J"])
    strata[key].append(row["row_id"])

rate = TARGET_N / len(csv_rows)
quotas = {}
remainders = []
for key, ids in strata.items():
    exact = len(ids) * rate
    floor = int(exact)
    quotas[key] = floor
    remainders.append((exact - floor, key))

need = TARGET_N - sum(quotas.values())
remainders.sort(key=lambda x: (-x[0], x[1]))
for i in range(need):
    _, key = remainders[i]
    quotas[key] += 1
assert sum(quotas.values()) == TARGET_N
for key in quotas:
    assert quotas[key] <= len(strata[key])

rng = random.Random(SEED)
sampled_ids = []
for key in sorted(strata.keys()):
    ids_sorted = sorted(strata[key], key=lambda x: int(x))
    k = quotas[key]
    if k > 0:
        sampled_ids.extend(rng.sample(ids_sorted, k))
# The original 71-unit stratified draw (seed 2026).
stratified_ids_set = set(sampled_ids)
assert len(sampled_ids) == TARGET_N

# --------------------------------------------------------------------------- #
# Full coverage of the arm under test. The stratified draw above lands only one
# of the 12 CP-GTR-originated units, which would leave agreement unmeasured on
# that arm. The rule is full coverage of CP-GTR's fresh units and proportional
# stratification elsewhere: the stratified draw is kept as drawn, and every
# CP-GTR-source unit not already in it is added. This is a deterministic union,
# not a second random draw, so no further seed is consumed.
# --------------------------------------------------------------------------- #
cp_gtr_full_ids_set = {row["row_id"] for row in csv_rows if row_source_arm[row["row_id"]] == "CP-GTR"}
assert len(cp_gtr_full_ids_set) == 12, len(cp_gtr_full_ids_set)
# Final overlap membership.
sampled_ids_set = stratified_ids_set | cp_gtr_full_ids_set
added_for_full_coverage = sorted(cp_gtr_full_ids_set - stratified_ids_set, key=int)


def balance(ids):
    """Counts by source arm, age and jurisdiction. Internal only: printed and
    recorded in overlap_sample.json, never written into either workbook."""
    sub_arm, sub_a, sub_j = defaultdict(int), defaultdict(int), defaultdict(int)
    for row in csv_rows:
        if row["row_id"] not in ids:
            continue
        sub_arm[row_source_arm[row["row_id"]]] += 1
        sub_a[row["A"]] += 1
        sub_j[row["J"]] += 1
    return dict(sub_arm), dict(sub_a), dict(sub_j)


full_arm, full_a, full_j = balance({row["row_id"] for row in csv_rows})
stratified_arm, stratified_a, stratified_j = balance(stratified_ids_set)
overlap_arm, overlap_a, overlap_j = balance(sampled_ids_set)

# --------------------------------------------------------------------------- #
# Practice units: the five practice pairs, deduplicated like the main 353. They
# are not part of the blind sheet and are never joined with any scored data.
# --------------------------------------------------------------------------- #
practice_pairs = json.loads(PRACTICE_PATH.read_text(encoding="utf-8"))
practice_scored = build_scored_documents(practice_pairs)
# Sorted (text_hash, A, J); the order is stable but arbitrary, since practice
# units carry no blindness requirement.
practice_units = unique_units(practice_scored)
practice_text_for_unit = {}
for d in practice_scored:
    practice_text_for_unit.setdefault(d["unit"], d["text"])

practice_rows = [
    {"unit_id": f"PRACTICE-{i}", "A": A, "J": J, "text": practice_text_for_unit[(th, A, J)]}
    for i, (th, A, J) in enumerate(practice_units)
]

# --------------------------------------------------------------------------- #
# Styling
# --------------------------------------------------------------------------- #
header_font = Font(name="Arial", size=11, bold=True, color=WHITE)
header_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
data_font = Font(name="Arial", size=10, bold=False)
data_align_wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)
data_align_nowrap = Alignment(horizontal="left", vertical="top", wrap_text=False)
data_align_center = Alignment(horizontal="center", vertical="top", wrap_text=False)
thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
ref_fill = PatternFill("solid", fgColor=NAVY)
judge_fill = PatternFill("solid", fgColor=ORANGE)
stripe_fill = PatternFill("solid", fgColor=STRIPE)

REF_COLS = [
    # (header, width, align)
    ("unit_id", 14, data_align_nowrap),
    ("reader_age", 11, data_align_center),
    ("jurisdiction", 12, data_align_center),
    ("document_text", 100, data_align_wrap),
]
CODE_COL_WIDTH = 13
NOTES_COL_WIDTH = 40
TOTAL_COLS = len(REF_COLS) + len(code_ids) + 1  # + notes


def code_comment(code_entry):
    aw = code_entry["applies_when"]
    lines = [
        code_entry["marking_criterion"],
        "",
        f"Applies when: jurisdiction={aw['jurisdiction']}, age_condition={aw['age_condition']}",
    ]
    if code_entry.get("citation"):
        lines.append(f"Citation: {code_entry['citation']}")
    return Comment("\n".join(lines), "codebook_v2")


def write_header(ws, header_row):
    for i, (header, width, _align) in enumerate(REF_COLS, start=1):
        col = get_column_letter(i)
        cell = ws[f"{col}{header_row}"]
        cell.value = header
        cell.font, cell.alignment, cell.fill = header_font, header_align, ref_fill
        cell.protection = Protection(locked=True)
        ws.column_dimensions[col].width = width
    for j, code_entry in enumerate(codes, start=1):
        i = len(REF_COLS) + j
        col = get_column_letter(i)
        cell = ws[f"{col}{header_row}"]
        cell.value = code_entry["code"]
        cell.font, cell.alignment, cell.fill = header_font, header_align, judge_fill
        cell.protection = Protection(locked=True)
        cell.comment = code_comment(code_entry)
        ws.column_dimensions[col].width = CODE_COL_WIDTH
    notes_col = get_column_letter(len(REF_COLS) + len(codes) + 1)
    cell = ws[f"{notes_col}{header_row}"]
    cell.value = "notes"
    cell.font, cell.alignment, cell.fill = header_font, header_align, judge_fill
    cell.protection = Protection(locked=True)
    ws.column_dimensions[notes_col].width = NOTES_COL_WIDTH
    ws.row_dimensions[header_row].height = 60


def write_rows(ws, rows, start_row, unlocked_counter):
    """rows: list of {"unit_id","A","J","text"}. Appends the number of
    applicable (unlocked) judgement cells to unlocked_counter, one entry per
    row written."""
    r_idx = start_row
    for stripe_i, row in enumerate(rows):
        A, J, text = row["A"], row["J"], row["text"]
        values = [row["unit_id"], A, J, text]
        stripe = stripe_fill if stripe_i % 2 == 1 else None
        max_lines = 1
        for c_idx, (val, (_h, width, align)) in enumerate(zip(values, REF_COLS), start=1):
            col = get_column_letter(c_idx)
            cell = ws[f"{col}{r_idx}"]
            cell.value = val
            cell.font, cell.alignment, cell.protection, cell.border = data_font, align, Protection(locked=True), border
            if stripe:
                cell.fill = stripe
            if c_idx == 4 and val:
                chars_per_line = max(10, width - 2)
                max_lines = max(max_lines, math.ceil(len(str(val)) / chars_per_line))

        unlocked_here = 0
        for j, code_entry in enumerate(codes, start=1):
            c_idx = len(REF_COLS) + j
            col = get_column_letter(c_idx)
            cell = ws[f"{col}{r_idx}"]
            cell.font, cell.alignment, cell.border = data_font, data_align_center, border
            if code_applies(code_entry, A, J):
                cell.value = None
                cell.protection = Protection(locked=False)
                unlocked_here += 1
            else:
                cell.value = "n-a"
                cell.protection = Protection(locked=True)
            if stripe:
                cell.fill = stripe
        unlocked_counter.append(unlocked_here)

        notes_col = get_column_letter(len(REF_COLS) + len(codes) + 1)
        cell = ws[f"{notes_col}{r_idx}"]
        cell.value = None
        cell.font, cell.alignment, cell.protection, cell.border = data_font, data_align_wrap, Protection(locked=False), border
        if stripe:
            cell.fill = stripe

        ws.row_dimensions[r_idx].height = max(40, max_lines * 14 + 10)
        r_idx += 1
    return r_idx


def add_dropdown_validation(ws, rows, start_row):
    """One DataValidation per code column, added as contiguous runs of
    unlocked (applicable) cells only, so locked n-a cells never get the
    yes/no dropdown."""
    for j, code_entry in enumerate(codes, start=1):
        c_idx = len(REF_COLS) + j
        col = get_column_letter(c_idx)
        dv = DataValidation(type="list", formula1='"yes,no"', allow_blank=True, showDropDown=False)
        dv.error = "Choose one of: yes, no"
        dv.errorTitle = "Invalid entry"
        dv.prompt = "yes / no"
        dv.promptTitle = f"{code_entry['code']}?"
        ws.add_data_validation(dv)
        run_start = None
        for offset, row in enumerate(rows):
            r_idx = start_row + offset
            applies = code_applies(code_entry, row["A"], row["J"])
            if applies and run_start is None:
                run_start = r_idx
            elif not applies and run_start is not None:
                dv.add(f"{col}{run_start}:{col}{r_idx - 1}")
                run_start = None
        if run_start is not None:
            dv.add(f"{col}{run_start}:{col}{start_row + len(rows) - 1}")


def sheet_protection():
    return SheetProtection(
        sheet=True, password=None,
        selectLockedCells=False, selectUnlockedCells=False,
        insertRows=True, insertHyperlinks=True, autoFilter=True,
        deleteColumns=True, insertColumns=True, pivotTables=True, deleteRows=True,
        formatCells=False, formatColumns=False, formatRows=False, sort=False,
    )


def build_practice_sheet(wb):
    """Banner on row 1 (merged, marking the sheet as practice-only), header on
    row 2, data from row 3."""
    ws = wb.create_sheet("Practice")
    ws.sheet_properties.tabColor = PRACTICE_BANNER_FILL

    last_col = get_column_letter(TOTAL_COLS)
    ws.merge_cells(f"A1:{last_col}1")
    banner = ws["A1"]
    banner.value = ("PRACTICE DATA -- for training only. Not scored, not part of the blind "
                     "annotation sheet, and never joined with annotator_primary.xlsx / annotator_overlap.xlsx.")
    banner.font = Font(name="Arial", size=11, bold=True, color=WHITE)
    banner.fill = PatternFill("solid", fgColor=PRACTICE_BANNER_FILL)
    banner.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 22

    write_header(ws, header_row=2)
    write_rows(ws, practice_rows, start_row=3, unlocked_counter=[])
    add_dropdown_validation(ws, practice_rows, start_row=3)
    ws.freeze_panes = "E3"
    ws.protection = sheet_protection()
    return ws


def build_annotation_sheet(wb, rows, sheet_name):
    ws = wb.create_sheet(sheet_name)
    write_header(ws, header_row=1)
    unlocked_counter = []
    write_rows(ws, rows, start_row=2, unlocked_counter=unlocked_counter)
    add_dropdown_validation(ws, rows, start_row=2)
    ws.freeze_panes = "E2"
    ws.protection = sheet_protection()
    return ws, unlocked_counter


def build_codes_sheet(wb):
    ws = wb.create_sheet("Codes")
    headers = ["code_id", "applies_when", "marking_criterion", "citation", "example_compliant", "example_noncompliant"]
    widths = [24, 40, 60, 24, 55, 55]
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        col = get_column_letter(i)
        cell = ws[f"{col}1"]
        cell.value = h
        cell.font, cell.alignment, cell.fill = header_font, header_align, ref_fill
        cell.protection = Protection(locked=True)
        ws.column_dimensions[col].width = w
    ws.row_dimensions[1].height = 22
    for r, code_entry in enumerate(codes, start=2):
        aw = code_entry["applies_when"]
        values = [
            code_entry["code"],
            f"jurisdiction={aw['jurisdiction']}, age_condition={aw['age_condition']}",
            code_entry["marking_criterion"],
            code_entry["citation"] or "(none)",
            code_entry["example_compliant"],
            code_entry["example_noncompliant"],
        ]
        max_lines = 1
        for c_idx, (val, w) in enumerate(zip(values, widths), start=1):
            col = get_column_letter(c_idx)
            cell = ws[f"{col}{r}"]
            cell.value = val
            cell.font = data_font
            cell.alignment = data_align_wrap
            cell.protection = Protection(locked=True)
            cell.border = border
            chars_per_line = max(10, w - 2)
            max_lines = max(max_lines, math.ceil(len(str(val)) / chars_per_line))
        ws.row_dimensions[r].height = max(30, max_lines * 14 + 10)
    ws.freeze_panes = "A2"
    ws.protection = SheetProtection(sheet=True, password=None, selectLockedCells=False, selectUnlockedCells=False)
    return ws


def build_about_sheet(wb, workbook_kind, n_rows, mean_unlocked):
    ws = wb.create_sheet("About")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100
    label_font = Font(name="Arial", size=10, bold=True)
    value_font = Font(name="Arial", size=10, bold=False)
    value_align = Alignment(horizontal="left", vertical="top", wrap_text=True)
    label_align = Alignment(horizontal="left", vertical="top")

    about_rows = [
        ("Workbook", workbook_kind),
        ("phase2_blind_annotation_sheet.csv SHA-256", csv_sha256),
        ("codebook_v2.json SHA-256 (whole file)", codebook_sha256_wholefile),
        ("codebook_v2.json embedded sha256 field", codebook_embedded_sha256 +
         " (basis: {codes, over_requirement} only -- see codebook_v2.json's own \"sha256\" field, "
         "not a whole-file hash; differs from the row above by design)"),
        ("Blind-sheet shuffle seed", "2026 (analysis/phase2_harness.py: build_blind_annotation_sheet())"),
        ("Overlap sample seed", "2026 (analysis/phase2_build_annotation_workbooks.py, stratified by "
                                 "source arm x reader_age x jurisdiction, largest-remainder allocation, "
                                 "independent reuse of the same seed value as the blind-sheet shuffle)"),
        ("Annotation sheet rows", str(n_rows)),
        ("Mean unlocked (applicable) judgements per row", f"{mean_unlocked:.2f} of 11 codes"),
        ("What a unit is", "A unit is a (document_text, reader_age, jurisdiction) triple. Labels attach "
                            "to UNITS, not to systems: multiple six-arm outputs across the Phase 2 "
                            "generation run can share the identical unit when their text and context "
                            "coincide, and are labelled once, blindly, with no record here of which "
                            "arm(s) produced them."),
    ]
    for i, (label, value) in enumerate(about_rows, start=1):
        a, b = ws[f"A{i}"], ws[f"B{i}"]
        a.value, b.value = label, value
        a.font, b.font = label_font, value_font
        a.alignment, b.alignment = label_align, value_align
        lines = math.ceil(len(str(value)) / 98) if value else 1
        ws.row_dimensions[i].height = max(18, lines * 14 + 8)
    return ws


def verify_no_leaked_columns(rows_source, allowed_keys):
    leaked = set()
    for row in rows_source:
        leaked |= (set(row.keys()) - allowed_keys)
    assert not leaked, f"leaked columns detected: {leaked}"


# CSV rows carry row_id, A, J, text, the (empty) code columns and notes; confirm
# no other key exists before using them.
verify_no_leaked_columns(csv_rows, {"row_id", "A", "J", "text", "notes"} | set(code_ids))

primary_rows = [{"unit_id": f"U{int(row['row_id']):03d}", "A": int(row["A"]), "J": row["J"], "text": row["text"]}
                 for row in csv_rows]
overlap_rows = [r for r, csvrow in zip(primary_rows, csv_rows) if csvrow["row_id"] in sampled_ids_set]
FINAL_N = len(sampled_ids_set)
assert len(overlap_rows) == FINAL_N == TARGET_N + len(added_for_full_coverage)

# --------------------------------------------------------------------------- #
# Build PRIMARY workbook (all 353 units)
# --------------------------------------------------------------------------- #
wb_primary = openpyxl.Workbook()
wb_primary.remove(wb_primary.active)
build_practice_sheet(wb_primary)
_, primary_unlocked = build_annotation_sheet(wb_primary, primary_rows, "Annotation")
build_codes_sheet(wb_primary)
primary_mean_unlocked = sum(primary_unlocked) / len(primary_unlocked)
build_about_sheet(wb_primary, "annotator_primary.xlsx -- all 353 units", len(primary_rows), primary_mean_unlocked)
wb_primary.save(PRIMARY_PATH)

# --------------------------------------------------------------------------- #
# Build OVERLAP workbook (stratified subset plus CP-GTR units; same CSV order
# and same unit_ids as the primary workbook)
# --------------------------------------------------------------------------- #
wb_overlap = openpyxl.Workbook()
wb_overlap.remove(wb_overlap.active)
build_practice_sheet(wb_overlap)
_, overlap_unlocked = build_annotation_sheet(wb_overlap, overlap_rows, "Annotation")
build_codes_sheet(wb_overlap)
overlap_mean_unlocked = sum(overlap_unlocked) / len(overlap_unlocked)
build_about_sheet(
    wb_overlap,
    f"annotator_overlap.xlsx -- stratified {TARGET_N}-unit sample UNION all {len(cp_gtr_full_ids_set)} "
    f"CP-GTR-originated units ({FINAL_N} units total, {100*FINAL_N/len(csv_rows):.1f}%). Rule: full "
    "coverage of CP-GTR's fresh units, proportional stratification elsewhere (2026-09-24b follow-up).",
    len(overlap_rows), overlap_mean_unlocked)
wb_overlap.save(OVERLAP_PATH)

# --------------------------------------------------------------------------- #
# overlap_sample.json records both the original stratified-only draw and the
# extended (union) membership actually used to build annotator_overlap.xlsx.
# --------------------------------------------------------------------------- #
overlap_sample_out = {
    "seed": SEED,
    "n_total": len(csv_rows),
    "rule": "Full coverage of the primary arm's (CP-GTR's) fresh units; proportional stratification "
            "elsewhere. The original 71-unit stratified draw is kept unchanged; every CP-GTR-source unit "
            "missing from it is added on top -- a deterministic rule-based union, not a second random draw.",
    "revision_history": [
        {
            "revision": "2026-09-24 (original)",
            "n_sampled": TARGET_N,
            "sample_rate": TARGET_N / len(csv_rows),
            "sampled_unit_ids": [f"U{int(rid):03d}" for rid in sorted(stratified_ids_set, key=int)],
            "stratified_by": "source arm (first-encountered, matches inheritance_report's own definition) x reader_age x jurisdiction",
            "allocation_method": "largest-remainder proportional allocation, per-stratum sampling via random.Random(2026).sample()",
            "balance": {"by_source_arm": stratified_arm, "by_age": stratified_a, "by_jurisdiction": stratified_j},
        },
        {
            "revision": "2026-09-24b (extended -- CURRENT, used to build annotator_overlap.xlsx)",
            "n_sampled": FINAL_N,
            "sample_rate": FINAL_N / len(csv_rows),
            "sampled_unit_ids": [f"U{int(rid):03d}" for rid in sorted(sampled_ids_set, key=int)],
            "added_for_full_cp_gtr_coverage": [f"U{int(rid):03d}" for rid in added_for_full_coverage],
            "n_cp_gtr_source_units_total": len(cp_gtr_full_ids_set),
            "reason": "The original stratified draw landed only 1 of the 12 CP-GTR-originated units, "
                      "leaving agreement unmeasured on the arm under test.",
            "balance": {"by_source_arm": overlap_arm, "by_age": overlap_a, "by_jurisdiction": overlap_j},
        },
    ],
    "balance": {
        "full_353": {"by_source_arm": full_arm, "by_age": full_a, "by_jurisdiction": full_j},
        "overlap_current": {"by_source_arm": overlap_arm, "by_age": overlap_a, "by_jurisdiction": overlap_j},
    },
    # Flat fields mirror the current (extended) membership, for readers that
    # expect a single non-versioned record.
    "n_sampled": FINAL_N,
    "sample_rate": FINAL_N / len(csv_rows),
    "sampled_unit_ids": [f"U{int(rid):03d}" for rid in sorted(sampled_ids_set, key=int)],
    "stratified_by": "source arm (first-encountered, matches inheritance_report's own definition) x reader_age x jurisdiction, PLUS forced full coverage of CP-GTR-source units",
    "allocation_method": "largest-remainder proportional allocation (base 71) + deterministic union with all CP-GTR-source units (2026-09-24b)",
}
OVERLAP_SAMPLE_JSON.write_text(json.dumps(overlap_sample_out, indent=2), encoding="utf-8")

# --------------------------------------------------------------------------- #
# Report
# --------------------------------------------------------------------------- #
print(f"Written: {PRIMARY_PATH}")
print(f"  Practice: {len(practice_rows)} rows, Annotation: {len(primary_rows)} rows x {TOTAL_COLS} cols, "
      f"Codes: {len(codes)} rows, About: present")
print(f"  Mean unlocked judgements/row: {primary_mean_unlocked:.3f}")
print()
print(f"Written: {OVERLAP_PATH}")
print(f"  Practice: {len(practice_rows)} rows, Annotation: {len(overlap_rows)} rows x {TOTAL_COLS} cols, "
      f"Codes: {len(codes)} rows, About: present")
print(f"  Mean unlocked judgements/row: {overlap_mean_unlocked:.3f}")
print()
print(f"Written: {OVERLAP_SAMPLE_JSON}")
print()
print(f"CP-GTR full-coverage extension: {len(added_for_full_coverage)} units added to the original "
      f"{TARGET_N}-unit stratified draw -> {FINAL_N} units total ({100*FINAL_N/len(csv_rows):.1f}%)")
print(f"  added: {[f'U{int(x):03d}' for x in added_for_full_coverage]}")
print()
print(f"Overlap sample balance (full 353 vs FINAL overlap {FINAL_N}), by source arm:")
for arm in sorted(full_arm):
    fpct = 100 * full_arm[arm] / len(csv_rows)
    opct = 100 * overlap_arm.get(arm, 0) / FINAL_N
    print(f"  {arm:28s} full={full_arm[arm]:3d} ({fpct:5.1f}%)  overlap={overlap_arm.get(arm,0):3d} ({opct:5.1f}%)")
print("by reader_age:")
for a in sorted(full_a, key=int):
    fpct = 100 * full_a[a] / len(csv_rows)
    opct = 100 * overlap_a.get(a, 0) / FINAL_N
    print(f"  {a:>3s}  full={full_a[a]:3d} ({fpct:5.1f}%)  overlap={overlap_a.get(a,0):3d} ({opct:5.1f}%)")
print("by jurisdiction:")
for j in sorted(full_j):
    fpct = 100 * full_j[j] / len(csv_rows)
    opct = 100 * overlap_j.get(j, 0) / FINAL_N
    print(f"  {j:>3s}  full={full_j[j]:3d} ({fpct:5.1f}%)  overlap={overlap_j.get(j,0):3d} ({opct:5.1f}%)")

print()
print("Source-column verification: CSV rows carry only "
      f"{sorted({'row_id','A','J','text','notes'} | set(code_ids))} -- "
      "no arm, library, prompt id, rule, flag, round-trip, or provenance column exists in the source "
      "data reaching either workbook. Dropped upstream by build_scored_documents()/"
      "build_blind_annotation_sheet() before the CSV was ever written: pair_id, arm, library_hash, "
      "prompt_id, flag_status, round_trip_result, rules_fired, provenance_verified, "
      "negation_filter_drops, negation_drops_available, host_parse_fallback, host_parse_attempts, "
      "host_n_nodes, host_n_edges, text_hash (as an exposed column), is_practice.")
