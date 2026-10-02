"""Task 2 of the 2026-09-25(e) follow-up (session_report_2026-09-25f.md):
guard-derivation verification workbook for the 13 unconditional LIB-FULL
rules, for Prof. Ali.

CRITICAL FINDING, surfaced here rather than acted on silently: the task's own
premise ("their defect is the missing guard, not the repair") does not hold
under evidence that already exists in this repo. docs/fidelity_audit_Ali_
filled2.csv -- Prof. Ali's own prior fidelity review of all 31 LIB-FULL rules
-- already covers exactly these 13 rules (matched by content_hash, all marked
phi_gated=no, confirming they are the unconditional set) and his verdict on
EVERY one of them is guard_correct=no, with a notes field stating the CITED
SOURCE PROVISION HAS NOTHING TO DO WITH THE RULE'S CLAIMED SUBJECT AT ALL --
not "right subject, missing threshold." Examples: a "data retention" rule
sourced from UN CRC text about military recruitment age; a "prohibit selling
child data" rule sourced from treaty-reservation law; a "parental consent"
rule sourced from the right to life (Article 6); a "consent withdrawal" rule
sourced from "privacy... at all stages of the [legal] proceedings."

Per this task's own instruction ("Do NOT invent guards where the provision
does not support one -- report those separately as underivable"): since none
of the 13 source provisions support their rule's claimed data-privacy
behavior AT ALL, none support a specific age/jurisdiction THRESHOLD either --
a guard cannot legitimately be derived from a provision that was never about
the rule's subject in the first place. Result: 0 of 13 derivable, 13 of 13
underivable, each for its own specific, evidenced reason (not a blanket
default).

This does not mean the workbook is pointless: Prof. Ali's fidelity review
answered "is this a faithful extraction" (no, for all 13); it did not answer
"given that, what should happen to these rules" (drop from the library
entirely vs. some other disposition) -- a DIFFERENT, still-open question this
workbook puts to him directly, alongside the literally-requested guard_correct
/suggested_guard columns (both trivially "no"/"none derivable" here, but kept
for format compliance and because the alternative was never asked for).
"""
import csv
import hashlib
import math
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.protection import SheetProtection
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parent.parent
FIDELITY_PATH = REPO_ROOT / "docs" / "fidelity_audit_Ali_filled2.csv"
OUT_PATH = REPO_ROOT / "annotation" / "guard_verification_workbook.xlsx"

NAVY, ORANGE, WHITE = "1F4E78", "E36C09", "FFFFFF"

UNCONDITIONAL_HASHES = [
    "955b5cb74d805d744270c9a1e2809e10c443f6f201442a41ce659126b176eee7",
    "75193b645534babe0b18cd68698513c3fdc2fba7c29256f968dded80814f0d74",
    "35bab113cb9fc920f27b6f60a577e661d95385df274f720a98c10218f799eaf3",
    "47443c1b6dee27afb3933639efbea88acce5cf32fc4d8e5a6d288bc058d9151f",
    "23e999bce96b7346ae1359eec1850737bc9fb8cfe30edcbcc35ae50d34b91a30",
    "bcdfac426f8fb808082c72afc0cb7952608fde3fde54e1ccaeb7e1df8e0e746b",
    "54499e4f23ed10b402583689b5f63e9b037f1d28f8e061c55e61b36427913eec",
    "33eb9ee1dd4c8d15e871fec30b9edaad6d7ee76746dece3f1bb30d80d7d8cec9",
    "7799508da7e362554f4578a36e4a9b7335ae848e5ef6fe9ad67281dcf8c2891e",
    "74ebc3fd78ee1e57b4fd086dd4f440fbfc66a2afddaef1ad41bdba0fb39cb487",
    "e3fcf2ccba772f3744d94670193af7ce926657d31529d9b1496ed44a274e1770",
    "9ef661e94b9947d5da7be1d4a0499795543e33570dedf49143063b250dbbabbf",
    "48d0e8fdba86c14b460cc0654b65b85e4c417c7f1c98fc7308ea395c1639e81a",
]

rows_by_hash = {r["rule_id"]: r for r in csv.DictReader(open(FIDELITY_PATH, encoding="utf-8"))}
assert all(h in rows_by_hash for h in UNCONDITIONAL_HASHES)
fidelity_sha256 = hashlib.sha256(FIDELITY_PATH.read_bytes()).hexdigest()

header_font = Font(name="Arial", size=11, bold=True, color=WHITE)
header_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
data_font = Font(name="Arial", size=10, bold=False)
data_align_wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)
data_align_nowrap = Alignment(horizontal="left", vertical="top", wrap_text=False)
thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
ref_fill = PatternFill("solid", fgColor=NAVY)
reviewer_fill = PatternFill("solid", fgColor=ORANGE)

COLUMNS = [
    # (header, key_or_None, width, wrap, locked, section)
    ("rule_id", "rule_id", 68, False, True, "ref"),
    ("rule_name", "rule_name", 45, True, True, "ref"),
    ("template", "obligation_template", 55, True, True, "ref"),
    ("source_file", "source_file", 16, False, True, "ref"),
    ("source_provision_text", "source_provision_text", 75, True, True, "ref"),
    ("prof_ali_prior_fidelity_note", "notes", 75, True, True, "ref"),
    ("proposed_guard", "proposed_guard", 40, True, True, "ref"),
    ("guard_derivation_status", "derivation_status", 22, False, True, "ref"),
    ("guard_correct", None, 16, False, False, "reviewer"),
    ("suggested_guard", None, 40, True, False, "reviewer"),
    ("recommended_disposition", None, 26, False, False, "reviewer"),
    ("reviewer_notes", None, 45, True, False, "reviewer"),
]

# All 13 -- proposed_guard is honestly "none derivable" for every one, each
# with its OWN reason (not a copy-pasted blanket line), per the instruction
# not to invent guards the source doesn't support.
rows_out = []
rule_names = {
    "955b5cb74d805d744270c9a1e2809e10c443f6f201442a41ce659126b176eee7": "RequireParentalConsentForChildProtection",
    "75193b645534babe0b18cd68698513c3fdc2fba7c29256f968dded80814f0d74": "Replace indefinite retention with bounded retention to comply with limitation on restrictions",
    "35bab113cb9fc920f27b6f60a577e661d95385df274f720a98c10218f799eaf3": "RequireParentalConsentForChildDataProcessing",
    "47443c1b6dee27afb3933639efbea88acce5cf32fc4d8e5a6d288bc058d9151f": "ReplaceIndefiniteRetentionWithBoundedDuration",
    "23e999bce96b7346ae1359eec1850737bc9fb8cfe30edcbcc35ae50d34b91a30": "AddParentalConsentRequirementForChildAdoption",
    "bcdfac426f8fb808082c72afc0cb7952608fde3fde54e1ccaeb7e1df8e0e746b": "RecruitmentAgeBoundRule",
    "54499e4f23ed10b402583689b5f63e9b037f1d28f8e061c55e61b36427913eec": "AddParentalConsentForChildConsentProcess",
    "33eb9ee1dd4c8d15e871fec30b9edaad6d7ee76746dece3f1bb30d80d7d8cec9": "Replace indefinite retention period with a bounded duration",
    "7799508da7e362554f4578a36e4a9b7335ae848e5ef6fe9ad67281dcf8c2891e": "InterpreterAssistanceRequirement",
    "74ebc3fd78ee1e57b4fd086dd4f440fbfc66a2afddaef1ad41bdba0fb39cb487": "AddWithdrawalPermissionToConsentProcess",
    "e3fcf2ccba772f3744d94670193af7ce926657d31529d9b1496ed44a274e1770": "Replace indefinite retention with a 30-day bounded period after instrument deposit",
    "9ef661e94b9947d5da7be1d4a0499795543e33570dedf49143063b250dbbabbf": "Provision 2",
    "48d0e8fdba86c14b460cc0654b65b85e4c417c7f1c98fc7308ea395c1639e81a": "Prohibit selling child data (UN)",
}

for h in UNCONDITIONAL_HASHES:
    r = rows_by_hash[h]
    reason = r["notes"].strip()
    out = dict(r)
    out["rule_id"] = h
    out["rule_name"] = rule_names[h]
    out["proposed_guard"] = f"NONE DERIVABLE -- source provision does not support the rule's claimed subject: {reason}"
    out["derivation_status"] = "underivable"
    rows_out.append(out)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Guard Verification"

for i, (header, _key, width, _wrap, _locked, section) in enumerate(COLUMNS, start=1):
    col = get_column_letter(i)
    cell = ws[f"{col}1"]
    cell.value = header
    cell.font, cell.alignment = header_font, header_align
    cell.fill = ref_fill if section == "ref" else reviewer_fill
    cell.protection = Protection(locked=True)
    ws.column_dimensions[col].width = width
ws.row_dimensions[1].height = 30

for r_idx, row in enumerate(rows_out, start=2):
    max_lines = 1
    for c_idx, (header, key, width, wrap, locked, section) in enumerate(COLUMNS, start=1):
        col = get_column_letter(c_idx)
        value = row.get(key) if key is not None else None
        cell = ws[f"{col}{r_idx}"]
        cell.value = value
        cell.font = data_font
        cell.alignment = data_align_wrap if wrap else data_align_nowrap
        cell.protection = Protection(locked=locked)
        cell.border = border
        if wrap and value:
            chars_per_line = max(10, width - 2)
            lines = math.ceil(len(str(value)) / chars_per_line)
            max_lines = max(max_lines, lines)
    ws.row_dimensions[r_idx].height = max(40, max_lines * 14 + 10)

ws.freeze_panes = "A2"

last_row = len(rows_out) + 1
guard_col = get_column_letter(next(i for i, c in enumerate(COLUMNS, 1) if c[0] == "guard_correct"))
dv1 = DataValidation(type="list", formula1='"yes,no"', allow_blank=True, showDropDown=False)
dv1.promptTitle = "Guard correct?"
dv1.prompt = "Is the proposed guard (if any) correct? All 13 rows here propose NO guard -- answer 'no' if you agree none can be derived, or use suggested_guard/recommended_disposition to say otherwise."
ws.add_data_validation(dv1)
dv1.add(f"{guard_col}2:{guard_col}{last_row}")

disp_col = get_column_letter(next(i for i, c in enumerate(COLUMNS, 1) if c[0] == "recommended_disposition"))
dv2 = DataValidation(type="list", formula1='"drop from LIB-FULL,keep unguarded (Ungated arm only, current status quo),guard anyway (state guard in suggested_guard),other (explain in reviewer_notes)"',
                     allow_blank=True, showDropDown=False)
dv2.promptTitle = "Recommended disposition"
dv2.prompt = "Given the source provision does not support this rule's claimed subject, what should happen to it?"
ws.add_data_validation(dv2)
dv2.add(f"{disp_col}2:{disp_col}{last_row}")

ws.protection = SheetProtection(
    sheet=True, password=None,
    selectLockedCells=False, selectUnlockedCells=False,
    insertRows=True, insertHyperlinks=True, autoFilter=True,
    deleteColumns=True, insertColumns=True, pivotTables=True, deleteRows=True,
    formatCells=False, formatColumns=False, formatRows=False, sort=False,
)

# =================================================================== About
about = wb.create_sheet("About")
about.column_dimensions["A"].width = 26
about.column_dimensions["B"].width = 100
label_font = Font(name="Arial", size=10, bold=True)
value_font = Font(name="Arial", size=10, bold=False)
value_align = Alignment(horizontal="left", vertical="top", wrap_text=True)
label_align = Alignment(horizontal="left", vertical="top")

about_rows = [
    ("Purpose", "Guard-derivation verification for the 13 unconditional (Phi always-True) rules in "
                "LIB-FULL -- Stage 2c of the 2026-09-25(e) instruction."),
    ("Headline finding", "0 of 13 rules have a derivable guard. This is NOT because a threshold was hard "
                          "to pick -- it's because each rule's own cited source provision (verified against "
                          "docs/fidelity_audit_Ali_filled2.csv, Prof. Ali's own prior fidelity review of all "
                          "31 LIB-FULL rules, matched here by content_hash) has NOTHING TO DO with the rule's "
                          "claimed data-privacy subject at all -- e.g. a 'data retention' rule sourced from UN "
                          "CRC text about military recruitment age, a 'prohibit selling data' rule sourced "
                          "from treaty-reservation law, a 'parental consent' rule sourced from the right to "
                          "life (Article 6). A guard cannot legitimately be derived from a provision that was "
                          "never about the rule's subject in the first place -- doing so anyway would be "
                          "inventing a justification, which this task's own instruction rules out."),
    ("What this means for the task's premise", "The instruction that prompted this workbook states 'their "
                          "defect is the missing guard, not the repair.' That premise does not hold for any "
                          "of these 13 under Prof. Ali's own already-collected evidence -- the defect is a "
                          "prior, more fundamental fidelity problem. Flagged here rather than silently "
                          "proceeding as if the premise held."),
    ("Why CP-GTR-Ungated still 'clears' some codes", "The fact that these rules' output text happens to "
                          "satisfy codebook_v2's gold criteria (e.g. EU_PARENTAL_CONSENT 16->6) does not "
                          "validate their fidelity -- codebook_v2 and the rule library were built "
                          "independently (codebook_v2.json's own 'source' field), so a generic, un-sourced "
                          "template like 'a parent or guardian must give consent' can coincidentally satisfy "
                          "a gold check for 'does the notice mention parental consent' without being a "
                          "faithful transcription of any real legal requirement its own citation claims."),
    ("Source of prof_ali_prior_fidelity_note column", "docs/fidelity_audit_Ali_filled2.csv, SHA-256 " + fidelity_sha256 +
                          " -- his guard_correct/notes columns, already on record, reused here rather than "
                          "asking him to re-answer a question he already answered."),
    ("What this workbook actually asks", "Not 'is this rule's source faithful' (already answered, no, for "
                          "all 13) -- but 'given that, what should happen to the rule': drop it from "
                          "LIB-FULL, keep it unguarded (current status quo for the Ungated arm), guard it "
                          "anyway despite the fidelity gap (state a guard in suggested_guard), or something "
                          "else (reviewer_notes). guard_correct/suggested_guard are kept for format "
                          "compliance with the literal request."),
    ("Rows", "13 -- exactly the unconditional (Phi always-True across cpgtr.admit.CONTEXTS's 9-cell grid) "
             "rules in LIB-FULL, verified by direct phi() evaluation, not inferred from source filename."),
    ("Not yet applied", "No guard has been applied to any rule. This workbook is sent for review only; "
                          "nothing in LIB-FULL or codebook_v3 changes until Prof. Ali's verdict returns."),
]
for i, (label, value) in enumerate(about_rows, start=1):
    a, b = about[f"A{i}"], about[f"B{i}"]
    a.value, b.value = label, value
    a.font, b.font = label_font, value_font
    a.alignment, b.alignment = label_align, value_align
    lines = math.ceil(len(str(value)) / 98) if value else 1
    about.row_dimensions[i].height = max(18, lines * 14 + 8)

wb.save(OUT_PATH)
print(f"Written: {OUT_PATH}")
print(f"Rows: {len(rows_out)} (all 13 unconditional rules)")
print(f"Derivable guards: 0 / 13")
print(f"Underivable: 13 / 13 (each with its own source-mismatch reason)")
print(f"fidelity_audit_Ali_filled2.csv SHA-256: {fidelity_sha256}")
